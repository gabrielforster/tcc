"""Loading the knowledge base the responsive agent answers from.

Formats the schedule names: PDF, DOCX, CSV, TXT (and Markdown, which is what the sample
corpus uses). PDF and DOCX need optional libraries; rather than adding two heavy
dependencies for a corpus that is mostly text, they are imported lazily and the loader
explains what to install if the file type actually appears.

A CSV becomes one document per row, because the tables in this domain -- instalment
options, fee schedules -- are looked up row-wise. Keeping the header in each row's text is
what makes a row retrievable on its own.
"""

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
SUPPORTED = TEXT_SUFFIXES | {".csv", ".pdf", ".docx"}


@dataclass
class Document:
    document_id: str
    text: str
    source: str
    metadata: dict[str, Any] = field(default_factory=dict)
    #: The file this came from, when one file yields several documents (CSV rows).
    #: Retrieval reports this level, because a customer is cited a document, not a row.
    parent_id: str | None = None

    @property
    def root_id(self) -> str:
        return self.parent_id or self.document_id


def load_directory(directory: Path | str) -> list[Document]:
    """Every supported file under a directory, sorted for reproducible ordering."""
    directory = Path(directory)
    if not directory.exists():
        raise FileNotFoundError(f"Knowledge base directory not found: {directory}")
    documents: list[Document] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            documents.extend(load_file(path, directory))
    return documents


def load_file(path: Path, root: Path | None = None) -> list[Document]:
    suffix = path.suffix.lower()
    identifier = str(path.relative_to(root)) if root else path.name
    metadata = {"source": str(path), "format": suffix.lstrip(".")}

    if suffix in TEXT_SUFFIXES:
        return [
            Document(
                document_id=identifier,
                text=path.read_text(encoding="utf-8"),
                source=str(path),
                metadata=metadata,
            )
        ]
    if suffix == ".csv":
        return _load_csv(path, identifier, metadata)
    if suffix == ".pdf":
        return [_load_pdf(path, identifier, metadata)]
    if suffix == ".docx":
        return [_load_docx(path, identifier, metadata)]
    raise ValueError(f"Unsupported file type: {path.suffix}")


def _load_csv(path: Path, identifier: str, metadata: dict) -> list[Document]:
    """One document per row, each carrying its column names so it reads standalone."""
    documents = []
    with path.open(encoding="utf-8", newline="") as handle:
        for index, row in enumerate(csv.DictReader(handle), start=1):
            text = "; ".join(f"{key}: {value}" for key, value in row.items() if value)
            documents.append(
                Document(
                    document_id=f"{identifier}:{index}",
                    text=text,
                    source=str(path),
                    metadata={**metadata, "row": index, "parent_id": identifier},
                    parent_id=identifier,
                )
            )
    return documents


def _load_pdf(path: Path, identifier: str, metadata: dict) -> Document:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ImportError(
            f"Reading {path.name} needs pypdf. Install the rag extra: uv sync --all-extras"
        ) from error
    reader = PdfReader(str(path))
    text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    return Document(
        document_id=identifier,
        text=text,
        source=str(path),
        metadata={**metadata, "pages": len(reader.pages)},
    )


def _load_docx(path: Path, identifier: str, metadata: dict) -> Document:
    try:
        import docx
    except ImportError as error:
        raise ImportError(
            f"Reading {path.name} needs python-docx. Install the rag extra: uv sync --all-extras"
        ) from error
    document = docx.Document(str(path))
    text = "\n\n".join(p.text for p in document.paragraphs if p.text.strip())
    return Document(document_id=identifier, text=text, source=str(path), metadata=metadata)
